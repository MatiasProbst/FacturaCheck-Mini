from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]


def run_app(files=None):
    # AppTest no necesita abrir un navegador. Se simula únicamente el selector;
    # la lectura real, validación y renderizado de tablas sí se ejecutan.
    files = files or {}
    with patch("streamlit.file_uploader", side_effect=lambda *args, **kwargs: files.get(kwargs["key"])):
        return AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)


def example(filename):
    return BytesIO((ROOT / "examples" / filename).read_bytes())


def test_initial_screen():
    app = run_app()
    assert not app.exception
    assert app.title[0].value == "FacturaCheck Mini"
    assert len(app.info) == 2
    assert len(app.dataframe) == 0


def test_both_files_display_independent_tables():
    app = run_app({
        "purchase_order": example("orden_compra.xlsx"),
        "supplier_invoice": example("factura_proveedor.xlsx"),
    })
    assert not app.exception
    assert len(app.dataframe) == 3
    assert len(app.success) == 2
    assert app.dataframe[0].value.iloc[0]["Precio unitario"] == "1.250,50"
    assert app.dataframe[1].value.iloc[0]["Precio unitario"] == "1.300,50"
    assert [metric.value for metric in app.metric] == ["4", "0", "4", "-3.225,50", "+1.925,00", "-5.150,50"]
    assert app.get("download_button")[-1].label == "Descargar informe Excel"
    assert app.dataframe[2].value.iloc[0]["Diferencia económica"] == "500,00"


def test_invalid_order_does_not_block_invoice():
    app = run_app({
        "purchase_order": BytesIO(b"archivo corrupto"),
        "supplier_invoice": example("factura_proveedor.xlsx"),
    })
    assert not app.exception
    assert len(app.error) == 1
    assert "Orden de compra: No se pudo procesar" in app.error[0].value
    assert len(app.dataframe) == 1
    assert not app.metric
    assert len(app.get("download_button")) == 2  # Solo ejemplos, nunca el informe.


def test_validation_errors_block_results():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["Producto", "Cantidad", "Precio unitario"])
    workbook.active.append(["A", -1, 10])
    stream = BytesIO()
    workbook.save(stream)
    app = run_app({"purchase_order": stream, "supplier_invoice": example("factura_proveedor.xlsx")})
    assert not app.exception
    assert "Cantidad" in app.error[0].value
    assert not app.metric
    assert len(app.dataframe) == 2


def test_one_file_does_not_compare():
    app = run_app({"purchase_order": example("orden_compra.xlsx")})
    assert not app.exception
    assert not app.metric
    assert len(app.dataframe) == 1


def test_download_receives_valid_report_matching_screen():
    import streamlit as st
    from openpyxl import load_workbook

    with patch("streamlit.download_button", wraps=st.download_button) as download:
        app = run_app({
            "purchase_order": example("orden_compra.xlsx"),
            "supplier_invoice": example("factura_proveedor.xlsx"),
        })
    assert not app.exception
    arguments = download.call_args.kwargs
    assert arguments["file_name"].endswith(".xlsx")
    book = load_workbook(BytesIO(arguments["data"]))
    assert book.sheetnames == ["Resumen", "Comparación", "Observaciones"]
    assert book["Resumen"]["B7"].value == -3225.5
    assert app.metric[3].value == "-3.225,50"
    assert book["Resumen"]["B5"].value == 1925
    assert app.metric[4].value == "+1.925,00"
    assert book["Resumen"]["B6"].value == -5150.5
    assert app.metric[5].value == "-5.150,50"


def test_show_comparison_regression_example_result_attributes():
    """Ejecutar la función real con Excel reales y comprobar el contrato completo."""
    from openpyxl import load_workbook
    import streamlit as st

    def render_examples():
        from pathlib import Path
        from app import show_comparison
        from excel_reader import read_excel

        root = Path.cwd()
        show_comparison(
            read_excel(root / "examples" / "orden_compra.xlsx"),
            read_excel(root / "examples" / "factura_proveedor.xlsx"),
        )

    with patch("streamlit.download_button", wraps=st.download_button) as download:
        app = AppTest.from_function(render_examples).run(timeout=20)
    assert not app.exception, [error.message for error in app.exception]
    metrics = {metric.label: metric.value for metric in app.metric}
    assert metrics["Importe facturado en exceso (+)"] == "+1.925,00"
    assert metrics["Importe facturado de menos (−)"] == "-5.150,50"
    assert metrics["Diferencia económica neta"] == "-3.225,50"
    assert len(app.dataframe[0].value) == 4
    book = load_workbook(BytesIO(download.call_args.kwargs["data"]))
    assert [book["Resumen"].cell(row, 2).value for row in (5, 6, 7)] == [1925, -5150.5, -3225.5]


def rerun_with_files(app, files):
    with patch("streamlit.file_uploader", side_effect=lambda *args, **kwargs: files.get(kwargs["key"])):
        return app.run(timeout=20)


def test_example_downloads_and_about_project():
    import streamlit as st

    with patch("streamlit.download_button", wraps=st.download_button) as download:
        app = run_app()
    assert not app.exception
    assert len(download.call_args_list) == 2
    for call, filename in zip(download.call_args_list, ["orden_compra.xlsx", "factura_proveedor.xlsx"]):
        assert call.kwargs["data"] == (ROOT / "examples" / filename).read_bytes()
        assert call.kwargs["file_name"] == filename
    assert "Acerca del proyecto" in [item.label for item in app.expander]
    assert any("no interpreta automáticamente facturas PDF" in item.value for item in app.markdown)


def test_state_filter_keeps_metrics_and_full_report():
    import streamlit as st
    from openpyxl import load_workbook

    files = {"purchase_order": example("orden_compra.xlsx"), "supplier_invoice": example("factura_proveedor.xlsx")}
    app = run_app(files)
    original_metrics = [metric.value for metric in app.metric]
    app.selectbox(key="state_filter").select("Diferencia de precio")
    with patch("streamlit.download_button", wraps=st.download_button) as download:
        app = rerun_with_files(app, files)
    assert not app.exception
    assert len(app.dataframe[-1].value) == 1
    assert app.dataframe[-1].value.iloc[0]["Producto"] == "Cuaderno ficticio A"
    assert [metric.value for metric in app.metric] == original_metrics
    report = load_workbook(BytesIO(download.call_args.kwargs["data"]))
    assert report["Comparación"].max_row == 5


def test_empty_filter_has_clear_message_and_all_columns():
    from comparison import COLUMNS

    files = {"purchase_order": example("orden_compra.xlsx"), "supplier_invoice": example("factura_proveedor.xlsx")}
    app = run_app(files)
    app.selectbox(key="state_filter").select("Correcto")
    app = rerun_with_files(app, files)
    assert not app.exception
    assert app.dataframe[-1].value.empty
    assert list(app.dataframe[-1].value.columns) == COLUMNS
    assert any("No hay productos que coincidan" in item.value for item in app.info)


def test_only_differences_toggle_and_filter_intersection():
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["Producto", "Cantidad", "Precio unitario"])
    book.active.append(["Cuaderno ficticio A", 10, 1250.50])
    book.active.append(["Lapicera ficticia B", 18, 350.25])
    stream = BytesIO()
    book.save(stream)
    files = {"purchase_order": example("orden_compra.xlsx"), "supplier_invoice": stream}
    app = run_app(files)
    assert len(app.dataframe[-1].value) == 3
    app.toggle(key="differences_only").set_value(True)
    app = rerun_with_files(app, files)
    assert not app.exception
    assert len(app.dataframe[-1].value) == 2
    assert not app.dataframe[-1].value["Estado"].str.contains("Correcto").any()
    app.selectbox(key="state_filter").select("Correcto")
    app = rerun_with_files(app, files)
    assert not app.exception
    assert app.dataframe[-1].value.empty
    app.toggle(key="differences_only").set_value(False)
    app = rerun_with_files(app, files)
    assert not app.exception
    assert len(app.dataframe[-1].value) == 1


def test_all_state_filters_preserve_input():
    import pandas as pd
    from app import ALL_STATES, STATE_LABELS, filter_results
    from comparison import COLUMNS

    table = pd.DataFrame([[str(i), None, None, None, None, 0, 0, 0, state]
                          for i, state in enumerate(STATE_LABELS)], columns=COLUMNS)
    original = table.copy(deep=True)
    for state in STATE_LABELS:
        result = filter_results(table, state, False)
        assert len(result) == 1
        assert result.iloc[0]["Estado"] == state
    assert len(filter_results(table, ALL_STATES, True)) == 5
    assert len(filter_results(table, ALL_STATES, False)) == 6
    pd.testing.assert_frame_equal(table, original)


def test_argentine_money_presentation_keeps_decimal_precision():
    from decimal import Decimal
    from app import format_money_ar, preview_price

    cases = {
        "0": "0,00", "1234.5": "1.234,50", "-5150.50": "-5.150,50",
        "1234567.895": "1.234.567,90", "-0.001": "0,00",
        "123456789012345678901234567890.125": "123.456.789.012.345.678.901.234.567.890,13",
    }
    for value, expected in cases.items():
        assert format_money_ar(Decimal(value)) == expected
    assert format_money_ar(None) == "—"
    assert preview_price("incorrecto") == "incorrecto"
