from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from excel_reader import ExcelValidationError, REQUIRED_COLUMNS, read_excel, validate_columns


def excel_bytes(headers, rows=()):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def test_reads_valid_file_and_preserves_data():
    stream = excel_bytes(REQUIRED_COLUMNS, [("Producto ficticio", 2, 12.50)])
    data = read_excel(stream)
    assert list(data.columns) == list(REQUIRED_COLUMNS)
    assert data.iloc[0].to_dict() == {
        "Producto": "Producto ficticio", "Cantidad": 2, "Precio unitario": 12.50,
    }
    # Releer el mismo archivo funciona aunque el cursor haya avanzado.
    pd.testing.assert_frame_equal(read_excel(stream), data)


def test_preserves_decimal_text_and_product_na_for_comparison():
    from decimal import Decimal
    from comparison import compare_documents

    exact_price = "0.12345678901234567890123456789"
    data = read_excel(excel_bytes(REQUIRED_COLUMNS, [("NA", 2, exact_price)]))
    assert data.loc[0, "Producto"] == "NA"
    assert data.loc[0, "Precio unitario"] == exact_price
    result = compare_documents(pd.DataFrame(columns=REQUIRED_COLUMNS), data)
    assert result.net_difference == Decimal("0.24691357802469135780246913578")


@pytest.mark.parametrize("missing", REQUIRED_COLUMNS)
def test_reports_each_missing_column(missing):
    stream = excel_bytes([column for column in REQUIRED_COLUMNS if column != missing])
    with pytest.raises(ExcelValidationError, match=f"Faltan columnas obligatorias: {missing}"):
        read_excel(stream)


def test_reports_all_missing_columns():
    with pytest.raises(ExcelValidationError, match="Producto, Cantidad, Precio unitario"):
        validate_columns(pd.DataFrame({"Otra": [1]}))


def test_allows_extra_columns_and_different_order():
    data = read_excel(excel_bytes(
        ["Notas", "Precio unitario", "Producto", "Cantidad"],
        [("Ficticio", 10, "A", 2)],
    ))
    assert data.loc[0, "Notas"] == "Ficticio"


@pytest.mark.parametrize("content", [b"", b"esto no es un Excel", b"PK\x03\x04incompleto"])
def test_unreadable_files_have_clear_error(content):
    with pytest.raises(ExcelValidationError, match="No se pudo procesar el archivo"):
        read_excel(BytesIO(content))


def test_rejects_other_extension():
    stream = excel_bytes(REQUIRED_COLUMNS)
    stream.name = "archivo.csv"
    with pytest.raises(ExcelValidationError, match="extensión .xlsx"):
        read_excel(stream)


def test_headers_without_rows_are_valid_structure():
    assert read_excel(excel_bytes(REQUIRED_COLUMNS)).empty


def test_completely_blank_sheet_is_invalid():
    with pytest.raises(ExcelValidationError, match="Faltan columnas"):
        read_excel(excel_bytes([]))


def test_reads_only_first_sheet():
    workbook = Workbook()
    workbook.active.append(list(REQUIRED_COLUMNS))
    workbook.active.append(["Primera", 1, 10])
    workbook.create_sheet("Otra").append(["Encabezado diferente"])
    stream = BytesIO()
    workbook.save(stream)
    assert read_excel(stream).loc[0, "Producto"] == "Primera"


@pytest.mark.parametrize("filename", ["orden_compra.xlsx", "factura_proveedor.xlsx"])
def test_example_files(filename):
    path = Path(__file__).resolve().parents[1] / "examples" / filename
    data = read_excel(path)
    assert list(data.columns) == list(REQUIRED_COLUMNS)
    assert len(data) == 3
