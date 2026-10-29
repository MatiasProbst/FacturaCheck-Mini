from decimal import Decimal
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from comparison import COLUMNS, compare_documents, format_money
from excel_reader import REQUIRED_COLUMNS
from report import create_report, observations


def document(*rows):
    return pd.DataFrame(rows, columns=REQUIRED_COLUMNS)


@pytest.fixture
def result():
    return compare_documents(
        document(("Correcto", 2, 10), ("Precio", 2, 10), ("Cantidad", 4, 10),
                 ("Ambas", 2, 10), ("Faltante", 1, 7)),
        document(("Correcto", 2, 10), ("Precio", 2, 12), ("Cantidad", 3, 10),
                 ("Ambas", 3, 8), ("Adicional", 2, 3)),
    )


def workbook(result):
    return load_workbook(BytesIO(create_report(result)), data_only=False)


def test_generates_xlsx_and_three_sheets(result):
    content = create_report(result)
    assert content.startswith(b"PK")
    assert workbook(result).sheetnames == ["Resumen", "Comparación", "Observaciones"]


def test_totals_positive_negative_and_net(result):
    assert result.excess == Decimal(14)
    assert result.shortfall == Decimal(-17)
    assert result.net_difference == Decimal(-3)
    assert result.requested_total == Decimal(107)
    assert result.billed_total == Decimal(104)
    assert result.net_difference == result.excess + result.shortfall
    assert result.net_difference == result.billed_total - result.requested_total


def test_summary_matches_engine(result):
    summary = workbook(result)["Resumen"]
    expected = [result.total, result.correct, result.different, result.excess,
                result.shortfall, result.net_difference, result.requested_total, result.billed_total]
    for index, value in enumerate(expected, start=2):
        exported = summary.cell(index, 2)
        assert exported.data_type == "n"
        assert Decimal(str(exported.value)) == Decimal(str(value))


def test_comparison_values_and_styles(result):
    sheet = workbook(result)["Comparación"]
    assert [cell.value for cell in sheet[1]] == COLUMNS
    for index, row in enumerate(result.table.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            exported = sheet.cell(index, column)
            if isinstance(value, Decimal):
                assert exported.data_type == "n"
                expected = Decimal(format_money(value)) if 4 <= column <= 8 else value
                assert Decimal(str(exported.value)) == expected
            else:
                assert exported.value == value
    assert sheet.freeze_panes == "B2"
    assert sheet.auto_filter.ref == "A1:I7"
    assert sheet["A1"].font.bold
    assert sheet["A1"].fill.fgColor.rgb == "00244563"
    assert sheet.column_dimensions["I"].width >= 38
    assert "0.00" in sheet["H2"].number_format
    assert sheet["I2"].fill.fgColor.rgb != sheet["I3"].fill.fgColor.rgb


def test_observations_cover_all_inconsistencies(result):
    entries = {entry[0]: entry for entry in observations(result)}
    assert len(entries) == result.different == 5
    assert "Correcto" not in entries
    assert "Precio unitario solicitado: 10; facturado: 12" in entries["Precio"][2]
    assert "Cantidad solicitada: 4; facturada: 3" in entries["Cantidad"][2]
    assert "Precio unitario" in entries["Ambas"][2]
    assert "Cantidad solicitada" in entries["Ambas"][2]
    assert "no en la factura" in entries["Faltante"][2]
    assert "no en la orden" in entries["Adicional"][2]
    assert entries["Faltante"][3] == Decimal(-7)
    assert entries["Adicional"][3] == Decimal(6)
    assert "Menor importe facturado" in entries["Cantidad"][2]
    assert "Mayor importe facturado" in entries["Precio"][2]
    assert workbook(result)["Observaciones"].max_row == 6


@pytest.mark.parametrize("rows", [[], [("A", 2, 10)]])
def test_no_differences_and_empty_documents(rows):
    result = compare_documents(document(*rows), document(*rows))
    assert result.excess == result.shortfall == result.net_difference == Decimal(0)
    assert observations(result) == []
    note = workbook(result)["Observaciones"]["C2"].value
    assert note == ("No se detectaron inconsistencias." if rows else "No hay productos para analizar.")


def test_report_does_not_mutate_result_or_recompute_differences(result):
    original = result.table.copy(deep=True)
    # El informe no necesita documentos originales ni calcula deltas desde cantidades/precios.
    result.table.loc[0, "Diferencia económica"] = Decimal("123.45")
    assert workbook(result)["Comparación"]["H2"].value == 123.45
    result.table.loc[0, "Diferencia económica"] = original.loc[0, "Diferencia económica"]
    create_report(result)
    pd.testing.assert_frame_equal(result.table, original)


def test_rounding_is_only_for_output_and_totals_are_exact():
    result = compare_documents(document(), document(("A", 1, "0.005"), ("B", 1, "0.005")))
    assert result.excess == result.billed_total == Decimal("0.010")
    book = workbook(result)
    assert book["Resumen"]["B5"].value == 0.01
    assert book["Comparación"]["H2"].value == 0.01
    assert "0.005" in book["Observaciones"]["C2"].value
    assert result.table.iloc[0]["Diferencia económica"] == Decimal("0.005")


def test_large_amounts_preserved_as_text_without_rounding_calculations():
    result = compare_documents(document(), document(("A", 2, "12345678901234567890.125")))
    assert result.excess == Decimal("24691357802469135780.250")
    cell = workbook(result)["Resumen"]["B5"]
    assert cell.data_type == "s"
    assert cell.value == "24691357802469135780.25"


def test_product_names_cannot_be_exported_as_formulas():
    result = compare_documents(document(), document(("=1+1", 1, 1)))
    book = workbook(result)
    assert book["Comparación"]["A2"].value == "=1+1"
    assert book["Comparación"]["A2"].data_type == "s"
    assert book["Observaciones"]["A2"].data_type == "s"
    assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)


def test_simultaneous_difference_can_have_zero_impact():
    result = compare_documents(document(("A", 2, 10)), document(("A", 4, 5)))
    assert result.different == 1
    assert result.excess == result.shortfall == Decimal(0)
    assert "Sin impacto económico neto" in observations(result)[0][2]


def test_integer_quantity_has_no_decimal_separator_and_fraction_is_preserved():
    result = compare_documents(
        document(("A", "10.0", 2), ("B", "1.25", 3)),
        document(("A", 10, 2), ("B", "1.25", 3)),
    )
    sheet = workbook(result)["Comparación"]
    for column in ("B", "C"):
        assert sheet[f"{column}2"].value == 10
        assert sheet[f"{column}2"].number_format == "0"
        assert sheet[f"{column}3"].value == 1.25
        assert sheet[f"{column}3"].number_format == "0.###############"
