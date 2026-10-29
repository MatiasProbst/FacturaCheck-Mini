"""Informe Excel en memoria a partir del resultado ya calculado por el motor."""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from comparison import (
    BOTH, COLUMNS, CORRECT, INVOICE_ONLY, ORDER_ONLY, PRICE, QUANTITY,
    ComparisonResult, format_money,
)


MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00;0.00'
SIGNED_FORMAT = '+#,##0.00;[Red]-#,##0.00;0.00'
QUANTITY_FORMAT = '0.###############'
STATE_COLORS = {
    CORRECT: "E2F0D9", PRICE: "FCE4D6", QUANTITY: "FFF2CC",
    BOTH: "F4CCCC", ORDER_ONLY: "D9EAF7", INVOICE_ONLY: "E4DFEC",
}


def observations(result: ComparisonResult) -> list[list]:
    """Una observación por inconsistencia; no vuelve a comparar ni calcular importes."""
    entries = []
    for row in result.table.to_dict("records"):
        state = row["Estado"]
        if state == CORRECT:
            continue
        details = []
        if state in (PRICE, BOTH):
            details.append(
                f"Precio unitario solicitado: {row['Precio solicitado']}; "
                f"facturado: {row['Precio facturado']}."
            )
        if state in (QUANTITY, BOTH):
            details.append(
                f"Cantidad solicitada: {row['Cantidad solicitada']}; "
                f"facturada: {row['Cantidad facturada']}."
            )
        if state == ORDER_ONLY:
            details.append("El producto figura en la orden, pero no en la factura. Importe facturado: cero.")
        if state == INVOICE_ONLY:
            details.append("El producto figura en la factura, pero no en la orden. Importe solicitado: cero.")
        delta = row["Diferencia económica"]
        direction = "Mayor importe facturado" if delta > 0 else (
            "Menor importe facturado" if delta < 0 else "Sin impacto económico neto"
        )
        details.append(f"{direction}. Diferencia exacta (facturado − solicitado): {delta}.")
        entries.append([row["Producto"], state, " ".join(details), delta])
    return entries


def write_value(cell, value, *, money=False, signed=False):
    """Exportar valores, nunca fórmulas; redondear dinero solo para presentación.

    Excel conserva aproximadamente 15 dígitos significativos. Para cifras mayores
    se escribe texto decimal, evitando perder centavos al serializar el informe.
    """
    if isinstance(value, Decimal):
        text = format_money(value) if money else str(value)
        number = Decimal(text)
        digits = len(number.as_tuple().digits)
        if digits > 15 or abs(number.adjusted()) > 300:
            cell.value = text
            cell.data_type = "s"
        else:
            cell.value = number
        quantity_format = "0" if number == number.to_integral_value() else QUANTITY_FORMAT
        cell.number_format = (SIGNED_FORMAT if signed else MONEY_FORMAT) if money else quantity_format
        cell.alignment = Alignment(horizontal="right", vertical="top")
    elif isinstance(value, str):
        cell.value = value
        cell.data_type = "s"  # Un producto que comienza con '=' sigue siendo texto.
    else:
        cell.value = value


def style_table(sheet, widths):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="244563")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="243746")
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal, vertical="top", wrap_text=True,
            )
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F6FA")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def create_report(result: ComparisonResult) -> bytes:
    """Crear las tres hojas sin leer originales, cambiar resultados ni usar disco."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    comparison = workbook.create_sheet("Comparación")
    notes = workbook.create_sheet("Observaciones")
    summary.append(["FacturaCheck Mini · Resumen", "Valor"])
    metrics = [
        ("Total de productos analizados", result.total),
        ("Productos correctos", result.correct),
        ("Productos con diferencias", result.different),
        ("Importe facturado en exceso", result.excess),
        ("Importe facturado de menos", result.shortfall),
        ("Diferencia económica neta", result.net_difference),
        ("Importe total solicitado", result.requested_total),
        ("Importe total facturado", result.billed_total),
    ]
    for index, (label, value) in enumerate(metrics, start=2):
        write_value(summary.cell(index, 1), label)
        write_value(summary.cell(index, 2), value, money=True, signed=5 <= index <= 7)
    style_table(summary, [48, 30])
    summary.auto_filter.ref = None
    for row_index in (5, 6, 7):
        summary.cell(row_index, 2).font = Font(
            bold=True, color={5: "9C3500", 6: "1F4E78", 7: "243746"}[row_index],
        )
    explanations = [
        "Exceso: suma de diferencias positivas. De menos: suma de diferencias negativas (con signo menos).",
        "Neto = facturado − solicitado = exceso + de menos. Todos los importes usan la misma moneda de origen.",
        "Ausentes: importe cero en el documento donde no aparecen; cantidad y precio en blanco.",
        "Importes presentados a dos decimales (ROUND_HALF_UP). Totales calculados antes de redondear; pueden diferir de sumar cifras mostradas.",
        "Cifras que exceden la precisión numérica de Excel se conservan como texto decimal. Observaciones incluye el impacto exacto.",
    ]
    for index, note in enumerate(explanations, start=12):
        summary.merge_cells(start_row=index, start_column=1, end_row=index, end_column=2)
        summary.cell(index, 1, note).alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[index].height = 44

    comparison.append(COLUMNS)
    for index, row in enumerate(result.table.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            write_value(comparison.cell(index, column), value, money=4 <= column <= 8, signed=column == 8)
    style_table(comparison, [32, 20, 20, 22, 22, 24, 24, 24, 38])
    for row in comparison.iter_rows(min_row=2):
        row[8].fill = PatternFill("solid", fgColor=STATE_COLORS[row[8].value])
        comparison.row_dimensions[row[0].row].height = 34

    notes.append(["Producto", "Tipo de diferencia", "Observación", "Impacto económico"])
    entries = observations(result)
    for index, entry in enumerate(entries, start=2):
        for column, value in enumerate(entry, start=1):
            write_value(notes.cell(index, column), value, money=column == 4, signed=column == 4)
    if not entries:
        notes.cell(2, 3, "No se detectaron inconsistencias." if result.total else "No hay productos para analizar.")
    style_table(notes, [32, 38, 95, 25])
    for index in range(2, notes.max_row + 1):
        notes.row_dimensions[index].height = 92
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
