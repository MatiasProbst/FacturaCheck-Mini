"""Generar dos Excel ficticios reproducibles con openpyxl."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


def create_example(filename: str, rows: list[tuple]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos ficticios"
    sheet.append(["Producto", "Cantidad", "Precio unitario"])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A81")
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 22
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "#,##0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    EXAMPLES.mkdir(exist_ok=True)
    workbook.save(EXAMPLES / filename)


if __name__ == "__main__":
    create_example("orden_compra.xlsx", [
        ("Cuaderno ficticio A", 10, 1250.50),
        ("Lapicera ficticia B", 20, 350.25),
        ("Carpeta ficticia C", 5, 890.00),
    ])
    create_example("factura_proveedor.xlsx", [
        ("Cuaderno ficticio A", 10, 1300.50),
        ("Lapicera ficticia B", 18, 350.25),
        ("Marcador ficticio D", 3, 475.00),
    ])
